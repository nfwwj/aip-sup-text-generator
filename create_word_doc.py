from docx import Document
from io import BytesIO
from docx.shared import Pt
from docx.text.paragraph import Paragraph
num2words1 = {1: 'One', 2: 'Two', 3: 'Three', 4: 'Four', 5: 'Five', \
            6: 'Six', 7: 'Seven', 8: 'Eight', 9: 'Nine', 10: 'Ten', \
            11: 'Eleven', 12: 'Twelve', 13: 'Thirteen', 14: 'Fourteen', \
            15: 'Fifteen', 16: 'Sixteen', 17: 'Seventeen', 18: 'Eighteen', 19: 'Nineteen'}

def edit_existing_doc(data):
    """
    Fills in the placeholders within the word document template with the extracted and processed data, and saves as a file.
    
    Used in For_AIP_IC.py when the AIP SUP word documents are downloaded.
    
    """
    try:
        doc = Document("AIP SUP - Paya Lebar Crane (XXXXXX-XX).docx")

        # File name needs to be saved as capital, and word doc text needs to put as plural if required.
        num = int(data.get("num_cranes", 1))
        crane_plural = "cranes" if num > 1 else "crane"
        crane_plural2 = "Cranes" if num > 1 else "Crane"
        print(data)
        # Map values to placeholders in the word document
        replacements = {
            "{{start}}": str(data.get('start(yymmddhhmmz)', '')),
            "{{end}}": str(data.get('end(yymmddhhmmz)', '')),
            "{{num_cranes}}": str(num2words1.get(num, str(num))).lower(),
            "{{crane_plural}}": str(crane_plural),
            "{{crane_plural2}}": str(crane_plural2),
            "{{crane_plural3}}": str(crane_plural2).upper(),
            "{{height}}": str(data.get('height(ft)', '')),
            "{{coordinates}}": str(data.get('rounded_coordinates', '')),
            "{{bearing}}": str(data.get('bearing', '')),
            "{{distance}}": str(data.get('distance', '')),
            "{{start_v2}}": str(data.get('start(ddmmmyyyy)', '')).upper(),
            "{{end_v2}}": str(data.get('end(ddmmmyyyy)', '')).upper(),
        }
        
        small_text_keys = ["{{start_v2}}", "{{end_v2}}"]

        def apply_to_paragraphs(paragraphs):
            """
            Logic to edit placeholders within the template word document. Try not to touch this code, as I have 0 idea how it works.
            """
            for p in paragraphs:
                for placeholder, value in replacements.items():
                    if placeholder in p.text:
                        found_in_run = False
                        for run in p.runs:
                            if placeholder in run.text:
                                run.text = run.text.replace(placeholder, value)

                                # Force font size 8
                                if placeholder in small_text_keys:
                                    run.font.size = Pt(8)
                                found_in_run = True
                        
                        if not found_in_run:
                            p.text = p.text.replace(placeholder, value)

        
        apply_to_paragraphs(doc.paragraphs)

        # Process Textboxes
        textboxes = doc.element.xpath('.//*[local-name()="txbxContent"]')
        for tbx in textboxes:
            for p_element in tbx.xpath('.//*[local-name()="p"]'):
                p = Paragraph(p_element, doc)
                apply_to_paragraphs([p])

        aip_start = str(data.get('start(yymmddhhmmz)')) 
        formatted_aip_start = aip_start[:6]

        app_num = data.get("application_num").upper()
        output_name = f"AIP SUP - Paya Lebar {crane_plural2} ({app_num}) EFF {formatted_aip_start}.docx"
        doc.save(output_name)
        
        return output_name 

    except Exception as e:
        return {"error": f"Something went wrong with creating the word doc file: {e}"}
